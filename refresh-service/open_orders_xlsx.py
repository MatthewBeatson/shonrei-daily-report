"""Total Product on Order -- open sales orders -> SKU Summary workbook.

Backs the "Total Product on Order" download on the production admin page
(production/admin). Pulls every Cin7 sale that is AUTHORISED but not
(fully) invoiced -- the SAME selection the daily report's "Sales on hand"
uses (cin7_order_statuses / cin7_invoice_statuses in reporting.settings,
order dated within the last 365 days, see daily_refresh_supabase.cin7_sales)
-- and builds an .xlsx laid out like Cin7's sales report export plus the
SKU Summary sheet from the sku-summary-segments skill:

  Sheet        one row per open order line: SKU, Product, Order #, Order
               date, Customer, Customer reference, Category, Invoice
               status, Quantity (still to invoice), Tax, Total
  SKU Summary  consolidated SKU / Quantity spill list, tie-out check, a
               Customer filter cell (E4), and input-driven segment blocks

READ ONLY against Cin7 (GET saleList / sale / product). One sale detail
per open order at ~1/sec (Cin7's 60 calls/min limit), so a cold build of
~150 orders takes 2-3 minutes -- far past gunicorn's 120s timeout. So it
runs in a background thread (same pattern as /refresh) and the caller
polls /reports/open-orders/status, then fetches /reports/open-orders/
download. Sale details are cached in memory keyed on the saleList
signature, so repeat builds only re-read orders that changed.
"""
from __future__ import annotations

import io
import os
import re
import threading
import time
import zipfile
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from daily_refresh_supabase import cin7_list_signature, get_conn, load_settings, split

BASE = 'https://inventory.dearsystems.com/ExternalApi/v2/'
LOOKBACK_DAYS = 365
NZ = ZoneInfo('Pacific/Auckland')


def norm(v):
    return str(v or '').strip().upper()


def num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


# ------------------------------------------------------------------
# Cin7 (read only)
# ------------------------------------------------------------------

def _session():
    s = requests.Session()
    s.headers.update({
        'api-auth-accountid': os.environ['CIN7_ACCOUNT_ID'],
        'api-auth-applicationkey': os.environ['CIN7_API_KEY'],
        'Content-Type': 'application/json',
    })
    return s


def cin7_get(session, path, params=None, retries=7):
    delays = (5, 10, 20, 30, 45, 60)
    last = ''
    for attempt in range(retries):
        try:
            r = session.get(BASE + path, params=params, timeout=(30, 90))
            if r.status_code == 503:  # rate limit (60 calls / minute)
                last = 'HTTP 503 rate limit'
                time.sleep(62)
                continue
            if r.ok and (r.text or '').strip():
                return r.json()
            last = f'HTTP {r.status_code}: {(r.text or "")[:300]}'
        except requests.exceptions.RequestException as exc:
            last = f'{type(exc).__name__}: {exc}'
        if attempt < retries - 1:
            time.sleep(delays[min(attempt, len(delays) - 1)])
    raise RuntimeError(f'Cin7 {path} failed: {last}')


def _status_filters():
    """Same settings rows the daily report's Sales on hand reads, so the
    two can never disagree about which orders are "on hand"."""
    order_wanted = {'ORDERED', 'BACKORDERED'}
    invoice_wanted = {'DRAFT', 'NOT AVAILABLE', 'NOT INVOICED', 'PARTIALLY INVOICED'}
    try:
        conn = get_conn()
        try:
            cfg = load_settings(conn)
        finally:
            conn.close()
        order_wanted = {norm(x) for x in split(cfg.get('cin7_order_statuses'))} or order_wanted
        invoice_wanted = {norm(x) for x in split(cfg.get('cin7_invoice_statuses'))} or invoice_wanted
    except Exception as exc:  # noqa: BLE001 -- fall back to the seeded defaults
        print(f'open_orders: could not read settings, using defaults: {exc}', flush=True)
    if 'NOT AVAILABLE' in invoice_wanted:
        invoice_wanted.add('NOT INVOICED')
    return order_wanted, invoice_wanted


def open_sales(session):
    order_wanted, invoice_wanted = _status_filters()
    cutoff = datetime.now(timezone.utc) - timedelta(days=LOOKBACK_DAYS)
    since = cutoff.strftime('%Y-%m-%dT%H:%M:%S.000')
    out, page = [], 1
    while True:
        data = cin7_get(session, 'saleList', {'Page': page, 'Limit': 1000, 'createdSince': since})
        items = data.get('SaleList') or []
        for s in items:
            raw_date = s.get('OrderDate') or s.get('Created')
            try:
                if raw_date and datetime.fromisoformat(str(raw_date).replace('Z', '+00:00')).date() < cutoff.date():
                    continue
            except ValueError:
                pass
            inv = norm(s.get('CombinedInvoiceStatus') or s.get('InvoiceStatus')) or 'NOT AVAILABLE'
            if (norm(s.get('Status')) in order_wanted and norm(s.get('OrderStatus')) == 'AUTHORISED'
                    and inv in invoice_wanted and s.get('SaleID')):
                s['_inv'] = inv
                out.append(s)
        if len(items) < 1000:
            break
        page += 1
        if page > 100:
            raise RuntimeError('Cin7 saleList pagination exceeded 100 pages.')
    return out


def product_categories(session):
    cats, page = {}, 1
    while True:
        data = cin7_get(session, 'product', {'Page': page, 'Limit': 1000})
        items = data.get('Products') or []
        for p in items:
            if p.get('SKU'):
                cats[str(p['SKU']).strip()] = p.get('Category') or ''
        if len(items) < 1000:
            break
        page += 1
        time.sleep(1.1)
    return cats


def open_lines_for_sale(listing, detail):
    """Order lines with the quantity still to invoice (ordered - invoiced)."""
    order = detail.get('Order') or {}
    invoices = detail.get('Invoices') or []
    if isinstance(invoices, dict):
        invoices = [invoices]
    invoiced = {}
    for doc in invoices:
        if isinstance(doc, dict) and norm(doc.get('Status')) in {'AUTHORISED', 'PAID'}:
            for l in doc.get('Lines') or []:
                sku = str(l.get('SKU') or '').strip()
                invoiced[sku] = invoiced.get(sku, 0.0) + num(l.get('Quantity'))

    raw_date = listing.get('OrderDate') or detail.get('SaleOrderDate') or ''
    try:
        order_date = datetime.fromisoformat(str(raw_date).replace('Z', '+00:00')).date()
    except ValueError:
        order_date = None

    rows = []
    for l in order.get('Lines') or []:
        sku = str(l.get('SKU') or '').strip()
        qty = num(l.get('Quantity'))
        if not sku or qty <= 0:
            continue
        used = min(qty, invoiced.get(sku, 0.0))  # invoiced qty consumes lines in order
        invoiced[sku] = invoiced.get(sku, 0.0) - used
        remaining = qty - used
        if remaining <= 0:
            continue
        share = remaining / qty
        rows.append({
            'sku': sku,
            'product': l.get('Name') or '',
            'order': listing.get('OrderNumber') or detail.get('Order', {}).get('SaleOrderNumber') or '',
            'date': order_date,
            'customer': listing.get('Customer') or detail.get('Customer') or '',
            'reference': listing.get('CustomerReference') or detail.get('CustomerReference') or '',
            'invoice_status': listing['_inv'],
            'qty': remaining,
            'tax': round(num(l.get('Tax')) * share, 4),
            'total': round(num(l.get('Total')) * share, 2),
        })
    return rows



# ------------------------------------------------------------------
# Workbook
# ------------------------------------------------------------------

HEAD_FILL = PatternFill('solid', fgColor='F0F0F0')
INPUT_FILL = PatternFill('solid', fgColor='FFFF00')
INPUT_FONT = Font(bold=True, color='0000FF')
BOLD = Font(bold=True)
NOTE = Font(italic=True, color='808080')


def build_workbook(rows, out_path, order_count, customer_filter=None, pulled_at=None):
    pulled_at = pulled_at or datetime.now()
    rows = sorted(rows, key=lambda r: (r['sku'], r['order']))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet'

    ws['A1'] = 'Report: Open sales orders - Authorised, not invoiced (same SOs as "Sales on hand")'
    ws['A2'] = f'Pulled from Cin7 Core: {pulled_at:%d-%b-%Y %H:%M}'
    ws['A3'] = (f'Customer: {customer_filter}' if customer_filter else 'Customer: All') + \
               f'  |  {order_count} orders  |  orders dated within last {LOOKBACK_DAYS} days'
    ws['A4'] = 'Quantity = quantity still to invoice (partially invoiced orders show the balance only)'
    for c in ('I5', 'J5', 'K5'):
        ws[c] = 'Grand Total'
    headers = ['SKU', 'Product', 'Order #', 'Order date', 'Customer', 'Customer reference',
               'Category', 'Invoice status', 'Quantity', 'Tax', 'Total']
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=6, column=i, value=h)
        cell.fill = HEAD_FILL
        cell.font = BOLD

    first = 7
    for n, r in enumerate(rows):
        row = first + n
        vals = [r['sku'], r['product'], r['order'], r['date'], r['customer'], r['reference'],
                r['category'], r['invoice_status'], r['qty'], r['tax'], r['total']]
        for c, v in enumerate(vals, 1):
            ws.cell(row=row, column=c, value=v)
        ws.cell(row=row, column=4).number_format = 'dd-mmm-yyyy'
        ws.cell(row=row, column=9).number_format = '#,##0.##'
        ws.cell(row=row, column=10).number_format = '#,##0.00'
        ws.cell(row=row, column=11).number_format = '#,##0.00'
    last = first + max(len(rows), 1) - 1

    widths = {'A': 18, 'B': 70, 'C': 11, 'D': 12, 'E': 30, 'F': 26, 'G': 22, 'H': 18, 'I': 10, 'J': 10, 'K': 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A7'
    ws.auto_filter.ref = f'A6:K{last}'

    # ---------------- SKU Summary ----------------
    sm = wb.create_sheet('SKU Summary')
    SKU = f'Sheet!$A${first}:$A${last}'
    QTY = f'Sheet!$I${first}:$I${last}'
    CUS = f'Sheet!$E${first}:$E${last}'
    CRIT = '"*"&$E$4&"*"'  # customer filter cell; blank = all customers

    def sumifs(spill):
        return f'SUMIFS({QTY},{SKU},{spill},{CUS},{CRIT})'

    sm['A1'] = 'SKU'
    sm['B1'] = 'Quantity'
    sm['A1'].font = sm['B1'].font = BOLD
    sm['A2'] = ArrayFormula('A2', f'=_xlfn._xlws.SORT(_xlfn.UNIQUE(_xlfn._xlws.FILTER({SKU},'
                                  f'ISNUMBER(SEARCH($E$4,{CUS})),"no match")))')
    sm['B2'] = ArrayFormula('B2', '=' + sumifs('_xlfn.ANCHORARRAY(A2)'))

    sm['D1'] = 'Check: source total'
    sm['E1'] = f'=SUMIFS({QTY},{CUS},{CRIT})'
    sm['D2'] = 'Check: summary total'
    sm['E2'] = '=SUM(_xlfn.ANCHORARRAY(B2))'
    sm['D4'] = 'Customer filter'
    sm['E4'] = customer_filter or None
    sm['E4'].fill = INPUT_FILL
    sm['E4'].font = INPUT_FONT
    sm['D5'] = 'blank = all customers; part of a name works (e.g. Prouds)'
    for c in ('D1', 'D2', 'D4'):
        sm[c].font = BOLD
    sm['D5'].font = NOTE

    A = '_xlfn.ANCHORARRAY($A$2)'

    def block(col, inputs, note, sku_formula, guard):
        c = col
        letters = [get_column_letter(c + i) for i in range(len(inputs))]
        for i, v in enumerate(inputs):
            cell = sm.cell(row=1, column=c + i, value=v)
            cell.fill = INPUT_FILL
            cell.font = INPUT_FONT
            cell.alignment = Alignment(horizontal='left')
        if note:
            nc = sm.cell(row=1, column=c + max(1, len(inputs)), value=note)
            nc.font = NOTE
        L, Q = letters[0], get_column_letter(c + 1)
        sm.cell(row=2, column=c, value='SKU').font = BOLD
        sm.cell(row=2, column=c + 1, value='Quantity').font = BOLD
        f = sku_formula.format(c=L, d=get_column_letter(c + 1), e=get_column_letter(c + 2))
        qty = sumifs(f'_xlfn.ANCHORARRAY({L}3)')
        if guard:
            f = f'IF({L}$1="","",{f})'
            qty = f'IF({L}$1="","",{qty})'
        sm[f'{L}3'] = ArrayFormula(f'{L}3', '=' + f)
        sm[f'{Q}3'] = ArrayFormula(f'{Q}3', '=' + qty)
        sm.column_dimensions[L].width = 18
        sm.column_dimensions[Q].width = 10

    contains = '_xlfn._xlws.SORT(_xlfn._xlws.FILTER(' + A + ',ISNUMBER(FIND({c}$1,' + A + ')),"no match"))'
    starts = '_xlfn._xlws.SORT(_xlfn._xlws.FILTER(' + A + ',LEFT(' + A + ',LEN({c}$1))={c}$1,"no match"))'
    vt_band = ('_xlfn._xlws.SORT(_xlfn._xlws.FILTER(' + A + ',(LEFT(' + A + ',2)="VT")*(IFERROR(--MID('
               + A + ',3,4),0)>{c}$1),"no match"))')
    p_digit = ('_xlfn._xlws.SORT(_xlfn._xlws.FILTER(' + A + ',(LEFT(' + A + ',1)="P")*ISNUMBER(--MID('
               + A + ',2,1)),"no match"))')
    v_digit = ('_xlfn._xlws.SORT(_xlfn._xlws.FILTER(' + A + ',(LEFT(' + A + ',1)="V")*ISNUMBER(--MID('
               + A + ',2,1)),"no match"))')
    three = ('_xlfn._xlws.SORT(_xlfn._xlws.FILTER(' + A + ',(LEFT(' + A + ',3)={c}$1)+(LEFT(' + A
             + ',3)={d}$1)+(LEFT(' + A + ',3)={e}$1),"no match"))')

    block(7, ['BU'], 'SKUs containing this fragment', contains, True)            # G
    block(10, ['PT200'], 'SKUs containing this fragment', contains, True)        # J
    block(13, ['VT1418'], 'SKUs containing this fragment', contains, True)       # M
    block(16, [None], 'Spare - type a fragment here', contains, True)            # P
    block(19, [1418], 'VT + 4 digits above threshold', vt_band, False)          # S
    block(22, ['P + digits'], 'e.g. P110, P150 (excludes PT)', p_digit, False)  # V
    block(25, ['MMS', 'MMW', 'MTS'], None, three, False)                         # Y
    block(29, ['V + digits'], 'e.g. V382, V886 (excludes VT)', v_digit, False)  # AC
    block(32, ['Q'], 'SKUs beginning with this prefix', starts, True)            # AF
    sm.column_dimensions['A'].width = 18
    sm.column_dimensions['D'].width = 20
    sm.column_dimensions['E'].width = 16
    sm.freeze_panes = 'A3'

    wb.active = 1
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    return out_path


_METADATA_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
    'xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray">'
    '<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" '
    'pasteAll="1" pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" '
    'clearComments="1" assign="1" coerce="1" cellMeta="1"/></metadataTypes>'
    '<futureMetadata name="XLDAPR" count="1"><bk><extLst><ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}">'
    '<xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/></ext></extLst></bk></futureMetadata>'
    '<cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata></metadata>'
)


def _mark_dynamic_arrays(data: bytes) -> bytes:
    """openpyxl writes array formulas as old-style Ctrl+Shift+Enter arrays,
    which Excel shows as ONE cell instead of spilling. Tag them as dynamic
    (spilling) arrays the way Excel itself does: a metadata.xml part plus
    cm="1" on each array-formula cell."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            part = zin.read(item.filename)
            if item.filename.startswith('xl/worksheets/sheet') and item.filename.endswith('.xml'):
                part = re.sub(r'<c ([^>]*?)>(<f t="array")', r'<c \1 cm="1">\2', part.decode('utf-8')).encode('utf-8')
            elif item.filename == '[Content_Types].xml':
                part = part.replace(b'</Types>', b'<Override PartName="/xl/metadata.xml" ContentType='
                                    b'"application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/></Types>')
            elif item.filename == 'xl/_rels/workbook.xml.rels':
                part = part.replace(b'</Relationships>', b'<Relationship Id="rIdMeta1" Type="http://schemas.'
                                    b'openxmlformats.org/officeDocument/2006/relationships/sheetMetadata" '
                                    b'Target="metadata.xml"/></Relationships>')
            zout.writestr(item, part)
        zout.writestr('xl/metadata.xml', _METADATA_XML)
    return out.getvalue()


# ------------------------------------------------------------------
# Background job (one at a time, in-memory -- single gunicorn worker)
# ------------------------------------------------------------------

_lock = threading.Lock()
_detail_cache: dict[str, tuple[str, dict]] = {}  # SaleID -> (signature, sale detail)
_job = {'status': 'idle', 'done': 0, 'total': 0, 'started_at': None, 'finished_at': None,
        'error': None, 'orders': 0, 'lines': 0, 'units': 0, 'filename': None}
_file: bytes | None = None


def job_state():
    with _lock:
        return dict(_job)


def get_file():
    with _lock:
        return _file, _job.get('filename')


def _set(**kw):
    with _lock:
        _job.update(kw)


def _build():
    global _file
    try:
        session = _session()
        sales = open_sales(session)
        _set(total=len(sales))
        cats = product_categories(session)
        rows, fetched = [], 0
        for i, s in enumerate(sales, 1):
            sid = str(s['SaleID'])
            sig = cin7_list_signature(s)
            cached = _detail_cache.get(sid)
            if cached and cached[0] == sig:
                detail = cached[1]
            else:
                if fetched:
                    time.sleep(1.1)  # stay under 60 calls/minute
                detail = cin7_get(session, 'sale', {'ID': sid})
                if isinstance(detail.get('Sale'), dict):
                    detail = detail['Sale']
                _detail_cache[sid] = (sig, detail)
                fetched += 1
            for r in open_lines_for_sale(s, detail):
                r['category'] = cats.get(r['sku'], '')
                rows.append(r)
            _set(done=i)
        now = datetime.now(NZ)
        buf = io.BytesIO()
        build_workbook(rows, buf, len(sales), pulled_at=now)
        data = _mark_dynamic_arrays(buf.getvalue())
        with _lock:
            _file = data
            _job.update(status='done', finished_at=now.isoformat(), orders=len(sales), lines=len(rows),
                        units=sum(r['qty'] for r in rows),
                        filename=f'Total_product_on_order_{now:%Y-%m-%d_%H%M}.xlsx')
        print(f'open_orders: built {len(sales)} orders / {len(rows)} lines ({fetched} details fetched)', flush=True)
    except Exception as exc:  # noqa: BLE001
        print(f'open_orders: build failed: {exc}', flush=True)
        _set(status='error', error=str(exc)[:500], finished_at=datetime.now(NZ).isoformat())


def start_build():
    """Starts a build unless one is already running. Returns the state."""
    with _lock:
        if _job['status'] == 'running':
            return dict(_job)
        _job.update(status='running', done=0, total=0, error=None,
                    started_at=datetime.now(NZ).isoformat(), finished_at=None)
    threading.Thread(target=_build, daemon=True).start()
    return job_state()
